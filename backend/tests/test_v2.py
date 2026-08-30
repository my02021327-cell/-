from __future__ import annotations

import hashlib
import json
import re
import uuid
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.main import app, datasets, models
from backend.app.services.derived_service import add_derived
from backend.app.services.dq_service import run_data_quality
from backend.app.services.excel_service import HEADERS, parse_imputation_log, parse_workbook
from backend.app.services import project_rules
from backend.app.services.health_service import (
    build_deltas,
    build_evidence,
    diagnose,
    evaluate_health,
    evaluate_operator_actions,
    evaluate_telemetry_signals,
)
from backend.app.services.project_rules import compute_diagnosis
from backend.app.services.recommendation_service import (
    build_recommendations,
    build_trend,
)


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = Path("/Users/JinU-Lim/Desktop/무제 폴더 4/outputs/01a03ce1-7173-7501-9079-a7734d294b80/BioGuard_standard_input_fixture.xlsx")
BUNDLE_PATH = ROOT / "artifacts/deployment/bioguard_a30_v2_h1_h7.joblib"


def workbook_bytes(
    start: date,
    days: int,
    *,
    reverse_rows: bool = False,
    invalid_last_feed: bool = False,
    trailing_blank_rows: int = 0,
    missing_actual_indices: set[int] | None = None,
    row_overrides: dict[int, dict[str, object]] | None = None,
) -> bytes:
    rows = []
    for index in range(days):
        current = start + timedelta(days=index)
        feed_a = 80.0 + index % 5
        feed_b = 82.0 + (index * 2) % 5
        if invalid_last_feed and index == days - 1:
            feed_a = -1.0
            feed_b = -1.0
        missing_actual = index in (missing_actual_indices or set())
        row = [
            current, feed_a, feed_b,
            None if missing_actual else 4700.0 + 11 * index,
            None if missing_actual else 4600.0 + 9 * index,
            None if missing_actual else 59.0 + (index % 4) * 0.2,
            None if missing_actual else 60.0 + (index % 3) * 0.2,
            37.0 + (index % 5) * 0.05, 37.2 + (index % 4) * 0.05,
            7.5 + (index % 4) * 0.01, 7.6 + (index % 3) * 0.01,
            2200.0 + index, 2250.0 + index,
            16000.0 + index * 2, 16200.0 + index * 2,
            3600.0 + index, 3650.0 + index,
            3.0 + (index % 4) * 0.01, 3.1 + (index % 3) * 0.01,
            2.1 + (index % 4) * 0.01, 2.2 + (index % 3) * 0.01,
            25000.0 + index * 3, 25200.0 + index * 3,
        ]
        for field, value in (row_overrides or {}).get(index, {}).items():
            row[HEADERS.index(field)] = value
        rows.append(row)
    if reverse_rows:
        rows.reverse()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BIOGUARD_INPUT"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for _ in range(trailing_blank_rows):
        sheet.append([None] * len(HEADERS))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload_bytes(client: TestClient, name: str, payload: bytes):
    return client.post(
        "/api/v1/datasets/upload",
        files={"file": (name, payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_bundle_load_and_contract():
    bundle = joblib.load(BUNDLE_PATH)
    assert bundle["system_id"] == "A30_INVERSE_ERROR_RELIABILITY_WEIGHT_V2"
    assert bundle["supported_horizons"] == [1, 2, 3, 4, 5, 6, 7]
    assert bundle["expert_pool"] == ["B10", "B11", "M2", "M1F", "RIDGE"]
    assert len(bundle["models"]) == 14
    for target in ("CH4_m3d_observed", "biogas_AB_m3d"):
        for horizon in range(1, 8):
            model = bundle["models"][f"{target}__h{horizon}"]
            assert model["ridge"]["features"] == ["feed_AB_tpd", "feed_A", "feed_B"]
            assert model["minimum_kernel_coverage"] == 0.90
    current = hashlib.sha256((ROOT / "outputs/11_final_model_lock.yaml").read_bytes()).hexdigest()
    assert bundle["provenance"]["phase11_lock_sha256"] == current


def test_new_horizons_have_independent_locked_specs_and_metrics():
    specs = json.loads((ROOT / "outputs/v2/v2_model_spec_lock.json").read_text())
    metrics = pd.read_csv(ROOT / "outputs/v2/v2_validation_metrics.csv", encoding="utf-8-sig")
    for target in ("CH4_m3d_observed", "biogas_AB_m3d"):
        for horizon in (2, 4, 5, 6):
            spec = specs[f"{target}__h{horizon}"]
            assert set(spec) == {"k", "tau_res", "ridge_alpha"}
            rows = metrics.loc[(metrics.target_name == target) & (metrics.horizon == horizon)]
            assert set(rows.period) == {"VALIDATION_2022", "RETROSPECTIVE_EVALUATION_2023"}
            assert rows[["RMSE", "MAE", "Bias", "R2_level", "R2_delta", "Skill_vs_B10"]].notna().all().all()


def test_excel_dq_raw_preservation_and_observed_ch4_definition():
    raw = parse_workbook(FIXTURE.read_bytes())
    assert list(raw.columns) == HEADERS + ["source_row"]
    clean, dq, summary = run_data_quality(raw)
    derived = add_derived(clean)
    assert summary == {
        "valid": 27, "missing": 38, "suspect": 1, "invalid": 0,
        "by_status": {"MISSING_ZERO_SENTINEL": 32, "VALID": 27, "MISSING": 6, "SUSPECT_CROSS_FIELD": 1},
        "raw_preserved": True, "suspect_values_retained": True,
    }
    first = derived.iloc[0]
    expected = first.biogas_A_m3d * first.ch4_purity_A_pct / 100 + first.biogas_B_m3d * first.ch4_purity_B_pct / 100
    assert np.isclose(first.CH4_m3d_observed, expected)
    assert derived.iloc[1:].CH4_m3d_observed.isna().all()
    assert dq.raw_value.notna().any()


def test_live_inference_availability_weights_and_purity():
    payload = workbook_bytes(date(2025, 1, 1), 30)
    descriptor = datasets.upload("valid_2025.xlsx", payload)
    frame = datasets.load_clean(descriptor["dataset_id"])
    expert, availability = models._expert_predictions(frame, "CH4_m3d_observed", 1)
    assert availability["current_target_observed"] is True
    assert expert["B10"] is not None and expert["M2"] is not None
    weights, audit = models._weights("CH4_m3d_observed", 1, frame.date.max(), availability["available"], True)
    assert np.isclose(sum(weights.values()), 1.0)
    assert audit["mode"] == "PAST_ONLY_RELIABILITY"
    forecast = models.forecast(descriptor["dataset_id"], persist=False)
    assert [row["horizon"] for row in forecast["methane"]["predictions"]] == list(range(1, 8))
    assert all(row["value"] >= 0 for row in forecast["methane"]["predictions"])
    assert forecast["purity_next_day"]["method"] == "DERIVED_FROM_CH4_AND_BIOGAS_FORECAST"
    assert "purity" not in forecast["methane"]


def test_minimum_api_contract_end_to_end():
    client = TestClient(app)
    assert client.get("/api/v1/health").json()["status"] == "CONNECTED"
    uploaded = upload_bytes(client, "api_2025.xlsx", workbook_bytes(date(2025, 6, 1), 30))
    assert uploaded.status_code == 200
    body = uploaded.json()
    dataset_id = body["dataset_id"]
    assert len(body["forecast"]["methane"]["predictions"]) == 7
    assert body["latest_data_date"] == "2025-06-30"
    assert body["forecast_origin_date"] == "2025-06-30"
    assert body["latest_usable_date"] == "2025-06-30"
    for suffix in ("", "/latest", "/history", "/forecast", "/data-quality"):
        response = client.get(f"/api/v1/datasets/{dataset_id}{suffix}")
        assert response.status_code == 200, (suffix, response.text)
    assert client.get(f"/datasets/{dataset_id}").status_code == 200


def test_upload_route_and_xlsx_contract_are_identical_front_to_back():
    adapter = (ROOT / "backend/static/bioguard_v2.js").read_text()
    html = (ROOT / "deliverables/yeongcheon_bioguard_multiview_final.html").read_text()
    upload_route = next(route for route in app.routes if route.path == "/api/v1/datasets/upload")
    assert "POST" in upload_route.methods
    assert "const UPLOAD_PATH = '/datasets/upload';" in adapter
    assert "uploadEndpoint:API+UPLOAD_PATH" in adapter
    assert "form.append('file',file)" in adapter
    assert "'Content-Type'" not in adapter
    assert 'id="localDataFileInput" accept=".xlsx"' in html


def test_backend_rejects_non_xlsx_extension_and_classifies_schema_errors():
    client = TestClient(app)
    wrong_type = upload_bytes(client, "not_supported.xls", workbook_bytes(date(2025, 1, 1), 30))
    assert wrong_type.status_code == 415
    assert wrong_type.json()["detail"]["code"] == "INVALID_FILE_TYPE"

    workbook = Workbook()
    workbook.active.title = "WRONG_SHEET"
    buffer = BytesIO()
    workbook.save(buffer)
    missing_sheet = upload_bytes(client, "missing_sheet.xlsx", buffer.getvalue())
    assert missing_sheet.status_code == 422
    assert missing_sheet.json()["detail"]["code"] == "BIOGUARD_INPUT_MISSING"


def test_short_history_keeps_actual_dataset_ready_but_forecast_unavailable():
    client = TestClient(app)
    uploaded = upload_bytes(client, "short.xlsx", workbook_bytes(date(2026, 1, 1), 5))
    assert uploaded.status_code == 200
    body = uploaded.json()
    assert body["status"] == "READY"
    assert body["forecast_status"] == "INSUFFICIENT_HISTORY"
    assert body["forecast"]["status"] == "INSUFFICIENT_HISTORY"
    assert body["forecast"]["methane"]["predictions"] == []
    history = client.get(f"/api/v1/datasets/{body['dataset_id']}/history").json()["rows"]
    assert len(history) == 5


def test_unsorted_rows_and_trailing_blanks_use_max_valid_date():
    payload = workbook_bytes(
        date(2025, 7, 1), 30, reverse_rows=True, trailing_blank_rows=4
    )
    parsed = parse_workbook(payload)
    assert parsed.date.is_monotonic_increasing
    assert len(parsed) == 30
    client = TestClient(app)
    body = upload_bytes(client, "unsorted.xlsx", payload).json()
    assert body["latest_data_date"] == "2025-07-30"
    assert body["forecast_origin_date"] == "2025-07-30"


def test_invalid_latest_model_inputs_step_back_forecast_origin():
    client = TestClient(app)
    payload = workbook_bytes(date(2025, 8, 1), 31, invalid_last_feed=True)
    body = upload_bytes(client, "invalid_latest.xlsx", payload).json()
    assert body["latest_data_date"] == "2025-08-31"
    assert body["forecast_origin_date"] == "2025-08-30"
    assert body["latest_usable_date"] == "2025-08-30"
    forecast = body["forecast"]
    assert forecast["origin_date"] == "2025-08-30"
    assert [item["target_date"] for item in forecast["methane"]["predictions"]] == [
        "2025-08-31", "2025-09-01", "2025-09-02", "2025-09-03",
        "2025-09-04", "2025-09-05", "2025-09-06",
    ]


def test_second_upload_and_dataset_url_reload_are_isolated_and_reproducible():
    client = TestClient(app)
    first = upload_bytes(client, "dataset_2025.xlsx", workbook_bytes(date(2025, 1, 1), 30)).json()
    second = upload_bytes(client, "dataset_2026.xlsx", workbook_bytes(date(2026, 2, 1), 30)).json()
    assert first["dataset_id"] != second["dataset_id"]
    first_rows = client.get(f"/api/v1/datasets/{first['dataset_id']}/history").json()["rows"]
    second_rows = client.get(f"/api/v1/datasets/{second['dataset_id']}/history").json()["rows"]
    assert all(row["date"].startswith("2025-") for row in first_rows)
    assert all(row["date"].startswith("2026-") for row in second_rows)
    restored = client.get(f"/api/v1/datasets/{second['dataset_id']}").json()
    assert restored["latest_data_date"] == second["latest_data_date"]
    again = client.get(f"/api/v1/datasets/{second['dataset_id']}/forecast").json()
    assert again == second["forecast"]


def test_operational_reliability_error_matures_only_with_truth():
    descriptor = datasets.upload(FIXTURE.name, FIXTURE.read_bytes())
    response = {
        "origin_date": "2023-09-14", "status": "AVAILABLE",
        "model": {"version": "TEST", "system_id": "TEST", "supported_horizons": [1]},
        "methane": {"predictions": [{"horizon": 1, "target_date": "2023-09-15", "value": 5000.0}]},
        "purity_next_day": None,
    }
    run_id = "test_" + uuid.uuid4().hex
    audit = {
        "forecasts": [{
            "target": "CH4_m3d_observed", "horizon": 1, "origin_date": "2023-09-14",
            "target_date": "2023-09-15", "expert_predictions": {"RIDGE": 5000.0},
        }],
        "api_response": response,
    }
    models.repository.save_forecast(run_id, descriptor["dataset_id"], response, audit)
    models.repository.mature_forecast_errors(datasets.load_clean(descriptor["dataset_id"]))
    errors = models.repository.reliability_errors("CH4_m3d_observed", 1, "2023-09-15")
    assert any(row["expert_id"] == "RIDGE" and np.isfinite(row["error"]) for row in errors)


def test_frontend_live_adapter_has_no_browser_inference_or_frozen_fallback():
    adapter = (ROOT / "backend/static/bioguard_v2.js").read_text()
    html = (ROOT / "deliverables/yeongcheon_bioguard_multiview_final.html").read_text()
    assert "window.BioGuardAPI" in adapter
    assert "XLSX." not in adapter
    assert ".fit(" not in adapter
    assert "FORECAST_DATA" not in adapter
    assert "HIST." not in adapter
    assert "INITIAL_RAW_VARIABLES" not in adapter
    assert "01_clean_base" not in adapter
    assert "bindBackendUpload" in adapter
    assert "handleBackendFileUpload" in adapter
    assert "uploadBioGuardDataset" in adapter
    assert "resetDatasetState('UPLOADING...')" in adapter
    assert "scheduleBackendRetry" in adapter
    assert "history.pushState" in adapter
    assert "getDataset" in adapter
    assert "buildYearButtons" in adapter
    assert "buildCh4RecorderSeries" in adapter
    assert "buildMissingBridges" in adapter
    assert "renderCh4TrendRecorder" in adapter
    assert "resolveSignalPresentation" in adapter
    # Same-origin only: no host, port or scheme may be baked into the console.
    assert '<script src="/static/bioguard_v2.js"></script>' in html
    assert "127.0.0.1" not in adapter and "localhost" not in adapter
    assert "http://" not in adapter and "https://" not in adapter
    assert "const API = '/api/v1';" in adapter
    assert "127.0.0.1" not in html.split("x-bioguard-legacy-reference")[0]
    assert "window.location.protocol !== 'file:'" in html
    assert "7-DAY CH4 FORECAST" in html
    assert "TREND RECORDER · CH-01 CH4 PRODUCTION" in html
    assert "ch4RecorderBridgeTrace" in html
    assert "ch4RecorderForecastUpper" not in html
    assert "95% CONF" not in html
    assert html.count("if (window.__BIOGUARD_BACKEND_ONLY__) return;") == 4
    assert 'id="sigDatePicker" class="hmi-date-input" aria-label="Select inspection date"' in html


def test_frontend_recorder_and_signal_matrix_contract():
    adapter = (ROOT / "backend/static/bioguard_v2.js").read_text()
    html = (ROOT / "deliverables/yeongcheon_bioguard_multiview_final.html").read_text()
    for element_id in (
        "ch4RecorderActualTrace", "ch4RecorderBridgeTrace",
        "ch4RecorderTransitionTrace", "ch4RecorderForecastTrace",
        "ch4RecorderActualMarkers", "ch4RecorderForecastMarkers",
        "ch4RecorderT0", "summaryLastActual", "summaryD1", "summaryD7",
        "summaryDelta7", "summaryT0", "telemetryUpdateLabel",
    ):
        assert f'id="{element_id}"' in html
    assert 'stroke="#00FF00"' in html
    assert 'stroke="#00FFFF"' in html
    assert 'stroke="#FF0000"' in html
    assert 'data-visual-bridge="true"' in adapter
    assert "gap>=1&&gap<=2" in adapter
    assert "dateEpoch(rows[right].date)" in adapter
    assert 'data-signal-status="' in adapter
    # The adapter must be a pure status -> presentation map: no frontend re-judgement
    # of DQ or health thresholds. The backend already resolves DQ-unusable to DATA.
    # Lamp + CSS only; the wording is the backend's status_text, so the project's
    # own vocabulary reaches the screen unchanged.
    assert "const SIGNAL_PRESENTATION=" in adapter
    assert "text:(signal&&signal.status_text)" in adapter
    assert "dqStatus===" not in adapter
    assert "dqStatus.startsWith" not in adapter
    assert "DISPLAY_BANDS" not in adapter
    assert "classifyDisplayVariable" not in adapter
    # Gauges read backend telemetry values; nothing is recomputed in the browser.
    assert "function backendSignal(variable)" in adapter
    assert "backendSignal(spec[0])" in adapter            # temperature + pH bay
    assert "V2.forecast.biogas.predictions" in adapter
    assert "const average" not in adapter

    # Four bays: D+1 CH4 forecast, D+1 purity forecast, total biogas, temp+pH.
    for renderer in ("renderForecastGauge", "renderPurityForecastGauge",
                     "renderBiogasForecastGauge", "renderTempPhGauge"):
        assert f"function {renderer}()" in adapter
    for element in ("gaugeForecastValue", "gaugePurityForecastValue",
                    "gaugeBiogasValue", "gaugeTempValue", "gaugePhValue",
                    "gaugeTempStatus", "gaugePhStatus"):
        assert f'id="{element}"' in html, element
    # The VFA/TA top gauge is gone; it lives in the matrix instead.
    assert 'id="gaugeVfaTaValue"' not in html.split("x-bioguard-legacy-reference")[0]

    # CURRENT is the dataset's forecast origin, never the OS date.
    assert "function isCurrentMode()" in adapter
    assert "V2.selectedIndex===V2.currentIndex" in adapter
    assert "new Date()" not in adapter and "Date.now" not in adapter

    # Gauges 1-3 switch mode with the selected date: D+1 predictions on the
    # forecast origin, that date's measured values on any earlier date. Each
    # branch may only read its own kind of source.
    forecast_bodies = "".join(
        adapter.split(f"function {name}()")[1].split("\n  }")[0]
        for name in ("renderForecastGauge", "renderPurityForecastGauge",
                     "renderBiogasForecastGauge")
    )
    assert "isCurrentMode()" in forecast_bodies
    assert "'FORECAST'" in forecast_bodies and "'ACTUAL'" in forecast_bodies
    assert "methane.predictions" in forecast_bodies
    assert "purity_next_day" in forecast_bodies
    assert "biogasForecastH1()" in forecast_bodies
    # The historical branch reads backend measurements, never a forecast.
    assert "CH4_m3d_observed" in forecast_bodies
    assert "backendSignal('ch4_purity')" in forecast_bodies
    assert "backendSignal('biogas_total')" in forecast_bodies
    # Historical titles say "실제", forecast titles say "예측".
    for title in ("메탄생성량 측정", "메탄순도 측정", "총 바이오가스 측정량",
                  "메탄생성량 예측", "메탄순도 예측", "총 바이오가스 생산량 예측"):
        assert title in adapter, title
    # A historical date reports its own readout, not T0's.
    assert "SELECTED ACTUAL: " in adapter
    assert "FORECAST ORIGIN: " in adapter

    # The matrix draws an explicit list, not a blind map over the payload.
    assert "const TELEMETRY_DISPLAY_ORDER=" in adapter
    assert "['VFA_TA','VFA','ALK','TAN','TS','VS','CODcr']" in adapter
    for hidden in ("temperature", "pH", "FAN", "ch4_purity",
                   "biogas_total", "ch4_observed"):
        assert f"'{hidden}'" not in adapter.split("TELEMETRY_DISPLAY_ORDER=")[1].split("]")[0]
    # The operator panel is rendered from the backend payload, never decided here.
    assert "function renderOperatorActions()" in adapter
    assert "V2.selectedRow&&V2.selectedRow.operator_actions" in adapter
    # The separate evidence panel, the annunciator detail line, the colour-policy
    # strip and the bridge/T0 legend entries were removed from the console.
    live = html.split("x-bioguard-legacy-reference")[0]
    for gone in ('id="evidenceBody"', "PROCESS STATUS EVIDENCE", 'id="annDetail"',
                 "tower-policy-strip", "DISPLAY BRIDGE · NOT MEASURED",
                 'class="legend-t0"'):
        assert gone not in live, gone
    assert "renderEvidence" not in adapter
    assert "evidenceBody" not in adapter
    # The DATA_REVIEW recommendation card is gone; data quality is not an
    # operating instruction.
    assert "DATA_REVIEW" not in adapter
    from backend.app.services import recommendation_service
    assert "DATA_REVIEW" not in recommendation_service.PRIORITY_ORDER
    # ACTUAL / FORECAST legend entries stay.
    assert "legend-actual" in live and "legend-forecast" in live
    # The card is icon + title + at most three lines; the rule paths, gates and
    # advisories a developer needs are kept but folded away.
    # At most three recommendation cards on screen; the rest fold away.
    assert "recs||[]).slice(0,3)" in adapter
    assert "op-details" in adapter and "상세보기" in adapter
    assert "payload.recommendations" in adapter
    # FAN is not surfaced anywhere in the operator console.
    assert "유리암모니아" not in adapter
    # Markers expose their backend value and provenance for verification.
    for attribute in ("data-date=", "data-value=", "data-source="):
        assert attribute in adapter, attribute
    assert "DERIVED_FROM_IMPUTED" in adapter
    # The annunciator carries the process severity alone; the data-confidence
    # line was removed from the console (it stays in the API payload).
    assert "'PROCESS: '+health.process_text" in adapter
    assert "setLamps(health.lamp" in adapter
    assert "DATA CONFIDENCE" not in adapter
    assert "if red" not in adapter and "heater" not in adapter.lower()
    assert 'id="operatorActionBody"' in html
    assert "등록된 운전자 조치 명령 없음" not in adapter
    # A burst of navigation clicks must step once per click, not once per settle.
    assert "function navigationCursor()" in adapter
    assert "V2.pendingIndex=bounded;" in adapter
    # Lamp classes stay in the frontend; the Korean status words now come from the
    # backend, which is where the project's own vocabulary lives.
    for css_class in ("t-green", "t-amber", "t-red", "t-gray", "signal-info"):
        assert css_class in adapter or css_class in html
    service = (ROOT / "backend/app/services/health_service.py").read_text()
    for korean in ("정상", "주의(통계)", "이상(통계)", "데이터확인",
                   "데이터정상", "데이터점검", "데이터없음", "조치필요"):
        assert korean in service, korean
    assert "UPDATE: DATASET EVENT" in adapter
    assert "setInterval(" not in adapter
    assert ".led-number.led-forecast { color: #00FFFF; }" in html
    assert ".led-number.led-unavailable { color: #808080; }" in html
    assert 'class="led-number led-unavailable" id="gaugeForecastValue"' in html
    # One helper drives all three mode-switching bays, so the forecast and the
    # actual can never be bound through different code paths.
    assert "function setModeGauge(" in adapter
    assert "setPrimaryGauge(id,finite(value)?value:null,digits,mode)" in adapter
    assert "function sanitizeNumeric(value)" in adapter
    assert "function sortUniqueRowsByDate(rows)" in adapter
    assert "FORECAST_REQUIRES_EXACTLY_7_FINITE_HORIZONS" in adapter
    assert "requestVersion" in adapter
    assert "data-chart-state" in adapter
    assert 'preserveAspectRatio="xMidYMid meet"' in html
    assert 'id="ch4RecorderMessage"' in html
    assert "vector-effect: non-scaling-stroke" in html
    # The trend recorder must not stretch its glyphs, and must draw real axes.
    # Every live chart keeps its aspect ratio so labels and markers cannot shear.
    # (The one remaining preserveAspectRatio="none" is inside a display:none
    #  LEGACY_REFERENCE_ONLY block that the adapter never touches.)
    for svg_id in ("ch4RecorderSvg", "trendSvg"):
        opening = re.search(rf'<svg id="{svg_id}"[^>]*>', html).group(0)
        assert 'preserveAspectRatio="xMidYMid meet"' in opening
        assert "aspect-ratio:" in opening
    assert "$('trendYTicks').innerHTML=" in adapter
    assert "$('trendXLabels').innerHTML=" in adapter
    assert "$('trendGrid').innerHTML=" in adapter
    # Global tower is backend process health only, including a GRAY data state.
    assert 'id="signalTower" data-tower-state="NONE"' in html
    assert '.signal-tower-3d[data-tower-state="DATA"] .lamp-bulb.red' in html
    assert ".annunciator-title.status-gray" in html and "#A0A0A0" in html
    assert "TOWER_TEXT_CLASS" in adapter


def test_local_html_origin_can_reach_backend_upload_contract():
    client = TestClient(app)
    preflight = client.options(
        "/api/v1/datasets/upload",
        headers={
            "Origin": "null",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert preflight.status_code == 200
    assert preflight.headers["access-control-allow-origin"] == "null"


def _variable(value, valid=True, **extra):
    return {"value": value, "valid": valid, **extra}


def _project_state(**overrides):
    """The base scenario from the project's own runSelfTest() in the shipped HTML."""
    base = {
        "temperature": {"A": _variable(38.5), "B": _variable(38.6)},
        "pH": {"A": _variable(7.9), "B": _variable(7.92)},
        "VFA": {"A": _variable(2900), "B": _variable(2950)},
        "ALK": {"A": _variable(17000), "B": _variable(17200)},
        "VFA_TA": {"A": _variable(0.17), "B": _variable(0.17)},
        "TAN": {"A": _variable(4000), "B": _variable(4050)},
        "FAN": {"A": _variable(400), "B": _variable(420)},
        "TS": {"A": _variable(2.6), "B": _variable(2.6)},
        "VS": {"A": _variable(1.1), "B": _variable(1.1)},
        "CODcr": {"A": _variable(25000), "B": _variable(25000)},
    }
    base.update(overrides)
    return base


# The 2023-09-17 snapshot recorded in section 8 of the specification and shipped
# as INITIAL_RAW_VARIABLES in the HTML.
PROJECT_SNAPSHOT = {
    "temperature": {"A": _variable(None, False, reason="SENSOR_FAULT"),
                    "B": _variable(None, False, reason="SENSOR_FAULT")},
    "pH": {"A": _variable(8.08), "B": _variable(8.1)},
    "VFA": {"A": _variable(2319.0), "B": _variable(2531.0)},
    "ALK": {"A": _variable(18428.0), "B": _variable(None, False, reason="CORRUPTED")},
    "VFA_TA": {"A": _variable(0.12584),
               "B": _variable(None, False, reason="UNAVAILABLE_ALK_B_INVALID")},
    "TAN": {"A": _variable(None, False, reason="STALE"), "B": _variable(None, False, reason="STALE")},
    "FAN": {"A": _variable(None, False, reason="UNAVAILABLE_STALE_TAN"),
            "B": _variable(None, False, reason="UNAVAILABLE_STALE_TAN")},
    "TS": {"A": _variable(2.88), "B": _variable(2.82)},
    "VS": {"A": _variable(1.11), "B": _variable(1.15)},
    "CODcr": {"A": _variable(26945.0), "B": _variable(22675.0)},
}


def test_migrated_engine_reproduces_the_projects_own_qa_scenarios():
    """The ten scenarios the project's runSelfTest() asserts, plus its two Excel
    QA cases. If the migration ever drifts from the shipped rules, these fail."""
    scenarios = [
        ("TEST1 normal -> NORMAL", _project_state(), None,
         lambda d: d["health"]["state"] == "NORMAL"),
        ("TEST2 one-sided RED -> DATA_SUSPECT, no Feed action",
         _project_state(VFA={"A": _variable(3700), "B": _variable(3000)}), None,
         lambda d: d["health"]["state"] == "DATA_SUSPECT" and d["commands"]["feedTarget"] is None),
        ("TEST3 VFA/TA>0.22 without trend evidence -> PROCESS_WATCH",
         _project_state(VFA_TA={"A": _variable(0.23), "B": _variable(0.24)}), None,
         lambda d: d["health"]["state"] == "PROCESS_WATCH"),
        ("TEST4 ACID EARLY -> PROCESS_ACTION, Feed = 0.75 x Q_ref",
         _project_state(VFA_TA={"A": _variable(0.25), "B": _variable(0.24)}),
         {"VFA_TA": {"A": 0.05, "B": 0.045}, "VFA": {"A": 750, "B": 720}},
         lambda d: d["health"]["state"] == "PROCESS_ACTION"
         and "ACID_EARLY" in d["health"]["code"] and d["commands"]["feedTarget"] == 135.0),
        ("TEST5 ACID SEVERE -> Feed = 0.55 x Q_ref",
         _project_state(VFA_TA={"A": _variable(0.5), "B": _variable(0.52)}), None,
         lambda d: d["health"]["state"] == "PROCESS_ACTION"
         and "ACID_SEVERE" in d["health"]["code"] and d["commands"]["feedTarget"] == 99.0),
        ("TEST6 one temperature channel invalid -> DATA_SUSPECT, thermal gate closed",
         _project_state(temperature={"A": _variable(0, False, reason="SENSOR_FAULT"),
                                     "B": _variable(38.5)}), None,
         lambda d: d["health"]["state"] == "DATA_SUSPECT"
         and d["commands"]["thermalGate"] == "CONDITION_NOT_VERIFIED"),
        ("TEST7 both valid temperatures <36 -> THERMAL ACTION, gate verified",
         _project_state(temperature={"A": _variable(35.0), "B": _variable(35.5)}), None,
         lambda d: d["health"]["state"] == "PROCESS_ACTION"
         and "THERMAL_ACTION_CANDIDATE" in d["health"]["code"]
         and d["commands"]["thermalGate"] == "CONDITION_VERIFIED"),
        ("TEST8 stale TAN -> FAN unavailable",
         _project_state(
             TAN={"A": _variable(0, False, reason="STALE", stale=True, sampleAgeDays=20),
                  "B": _variable(0, False, reason="STALE", stale=True, sampleAgeDays=20)},
             FAN={"A": _variable(400), "B": _variable(420)}), None,
         lambda d: d["classified"]["FAN"]["A"]["status"] == "INVALID"
         and d["classified"]["FAN"]["B"]["status"] == "INVALID"),
        ("TEST9 ALK low alone -> no NaHCO3 action",
         _project_state(ALK={"A": _variable(14000), "B": _variable(14200)}), None,
         lambda d: d["commands"]["alkaliGate"] == "CONDITION_NOT_VERIFIED"),
        ("TEST10 2023-09-17 snapshot -> NORMAL_DEGRADED_DATA / MEDIUM",
         PROJECT_SNAPSHOT, None,
         lambda d: d["health"]["state"] == "NORMAL_DEGRADED_DATA"
         and d["data_confidence"] == "MEDIUM"),
        ("E1 missing CODcr -> INVALID, no PROCESS_ACTION from that alone",
         _project_state(CODcr={"A": _variable(0, False, reason="MISSING"),
                               "B": _variable(25000)}), None,
         lambda d: d["classified"]["CODcr"]["A"]["status"] == "INVALID"
         and d["health"]["state"] != "PROCESS_ACTION"),
        ("E2 corrupted ALK_B -> ALK.B and VFA_TA.B both INVALID",
         PROJECT_SNAPSHOT, None,
         lambda d: d["classified"]["ALK"]["B"]["status"] == "INVALID"
         and d["classified"]["VFA_TA"]["B"]["status"] == "INVALID"),
    ]
    failures = []
    for name, raw_variables, deltas, expectation in scenarios:
        diagnosis = compute_diagnosis(raw_variables, deltas, None)
        if not expectation(diagnosis):
            failures.append(f"{name}: state={diagnosis['health']['state']} "
                            f"code={diagnosis['health']['code']} "
                            f"feed={diagnosis['commands']['feedTarget']}")
    assert not failures, failures


def test_thresholds_match_the_project_specification_verbatim():
    """The numbers themselves, so a silent edit anywhere fails loudly."""
    assert project_rules.THRESHOLDS["temperature"] == dict(
        lowRed=37.7, lowGreen=37.9, highGreen=40.0, highRed=40.2, unit="°C")
    assert project_rules.THRESHOLDS["pH"] == dict(
        lowRed=7.56, lowGreen=7.70, highGreen=8.14, highRed=8.21, unit="")
    assert project_rules.THRESHOLDS["VFA"] == dict(
        lowRed=2337, lowGreen=2428, highGreen=3378, highRed=3625, unit="mg/L")
    assert project_rules.THRESHOLDS["ALK"] == dict(
        lowRed=15008, lowGreen=15324, highGreen=18305, highRed=18629, unit="mg/L")
    assert project_rules.THRESHOLDS["VFA_TA"] == dict(
        lowRed=0.130, lowGreen=0.137, highGreen=0.207, highRed=0.215, unit="")
    assert project_rules.THRESHOLDS["TAN"] == dict(
        lowRed=2077, lowGreen=2731, highGreen=5634, highRed=5920, unit="mg/L")
    assert project_rules.THRESHOLDS["FAN"] == dict(
        lowRed=170, lowGreen=226, highGreen=716, highRed=820, unit="mg/L")
    assert project_rules.THRESHOLDS["TS"] == dict(
        lowRed=2.09, lowGreen=2.26, highGreen=2.94, highRed=3.16, unit="%")
    assert project_rules.THRESHOLDS["VS"] == dict(
        lowRed=0.77, lowGreen=0.87, highGreen=1.41, highRed=1.59, unit="%")
    assert project_rules.THRESHOLDS["CODcr"] == dict(
        lowRed=15614, lowGreen=18739, highGreen=36311, highRed=40152, unit="mg/L")
    assert project_rules.ACID == dict(
        watchVFA_TA=0.215, p95VFA_TA=0.207, earlyDeltaVFA_TA=0.04, earlyDeltaVFA=700,
        earlyFeedFactor=0.75, severeVFA_TA=0.45, severeFeedFactor=0.55)
    assert project_rules.BUFFER == dict(alkLow=15000, phLow=7.70, vfaTaTrigger=0.22)
    assert project_rules.AMMONIA == dict(
        watchFanLow=500, watchFanHigh=700, actionFan=700, vfaTaTrigger=0.22)
    assert project_rules.THERMAL == dict(lowC=36, highC=41)
    assert project_rules.CONCORDANCE["temperature"] == dict(verify=2.0, strong=2.5)
    assert project_rules.CONCORDANCE["VFA_TA"] == dict(verify=0.04)
    assert project_rules.TAN_MAX_AGE_DAYS == 14
    assert project_rules.Q_REF_TPD == 180.0
    # Specification section 5: the tower shows the final state, and every state
    # in the hierarchy has a colour.
    assert project_rules.STATE_TO_LAMP == {
        "NORMAL": "GREEN", "NORMAL_DEGRADED_DATA": "AMBER", "DATA_SUSPECT": "AMBER",
        "PROCESS_WATCH": "AMBER", "PROCESS_ACTION": "RED", "RECOVERY": "AMBER",
    }


def test_no_process_threshold_lives_outside_the_project_registry():
    """The registry is the only place a band may appear."""
    adapter = (ROOT / "backend/static/bioguard_v2.js").read_text()
    service = (ROOT / "backend/app/services/health_service.py").read_text()
    for number in ("37.9", "40.0", "7.70", "8.14", "2428", "3378", "15324", "18305",
                   "0.137", "0.207", "0.215", "0.45", "2731", "5634", "2.26", "2.94",
                   "0.87", "1.41", "18739", "36311", "15000", "0.22", "700", "180"):
        assert number not in adapter, f"threshold {number} leaked into the frontend"
        assert number not in service, f"threshold {number} leaked into health_service"
    assert "THRESHOLDS" not in adapter and "HEALTH_RULES" not in adapter


def test_every_project_variable_is_classified_not_dropped_to_info():
    """pH, VFA, ALK, TAN, TS, VS and CODcr all have project rules, so none of them
    may fall through to an informational row."""
    fields = [header for header in HEADERS if header != "date"]
    values = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 2900.0 / 17000.0, "VFA_TA_B": 2950.0 / 17200.0,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
        "ch4_purity_A_pct": 62.0, "ch4_purity_B_pct": 63.0,
        "biogas_A_m3d": 5000.0, "biogas_B_m3d": 4900.0,
        "biogas_AB_m3d": 9900.0, "CH4_m3d_observed": 6187.0,
        "feed_A_tpd": 90.0, "feed_B_tpd": 90.0,
    }
    valid = {field: "VALID" for field in fields}
    diagnosis = diagnose(values, valid)
    signals = {item["variable"]: item
               for item in evaluate_telemetry_signals(diagnosis, values, valid)}

    for variable in ("temperature", "pH", "VFA", "ALK", "VFA_TA", "TAN", "TS", "VS", "CODcr"):
        signal = signals[variable]
        assert signal["status"] == "NORMAL", (variable, signal["status"])
        assert signal["status_text"] == "정상"
        assert signal["rule_source"] == project_rules.RULE_SOURCE
        assert signal["rule_id"] == variable
        assert signal["reason"] == "WITHIN_GREEN_RANGE"
        assert signal["semantics"] == "STATISTICAL_RARITY_NOT_PROCESS_ACTION"
        assert signal["green_range"] == [
            project_rules.THRESHOLDS[variable]["lowGreen"],
            project_rules.THRESHOLDS[variable]["highGreen"],
        ]
    # FAN has no formula anywhere in the project and is never fabricated.
    assert signals["FAN"]["status"] == "DATA"
    assert signals["FAN"]["reason"] == project_rules.FAN_UNAVAILABLE_REASON
    # Purity / biogas / observed CH4 are the project's EXCEL_INFORMATIONAL set.
    for variable in ("ch4_purity", "biogas_total", "ch4_observed"):
        signal = signals[variable]
        assert signal["status"] == "DATA_NORMAL"
        assert signal["status_text"] == "데이터정상"
        assert signal["rule_source"] == "PROJECT_EXCEL_INFORMATIONAL"
        assert signal["reason"] == "NO_PROJECT_HEALTH_RULE_DEFINED_FOR_THIS_VARIABLE"

    health = evaluate_health(diagnosis)
    # Tower is the engine's final state, never a copy of a single variable.
    assert health["state"] == "NORMAL_DEGRADED_DATA"      # FAN keeps confidence below HIGH
    # A data-confidence downgrade alone must not raise the process lamp; the
    # project's combined mapping is still reported for audit.
    assert health["lamp"] == "GREEN"
    assert health["combined_lamp"] == "AMBER"
    assert health["data_review"] is True
    assert health["rule_id"] == "evaluateProcessHealth"


def test_statistical_colour_is_reported_with_its_tail_and_is_not_a_process_action():
    """Specification section 2: RED on one variable is rarity, not an action."""
    fields = [header for header in HEADERS if header != "date"]
    values = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2000.0, "VFA_B_mgL": 2010.0,          # both LOW-side RED
        "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1176, "VFA_TA_B": 0.1169,            # both LOW-side RED
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    valid = {field: "VALID" for field in fields}
    diagnosis = diagnose(values, valid)
    signals = {item["variable"]: item
               for item in evaluate_telemetry_signals(diagnosis, values, valid)}
    assert signals["VFA"]["status"] == "ACTION"
    assert signals["VFA"]["tail"] == {"A": "LOW", "B": "LOW"}
    assert signals["VFA"]["reason"] == "RED_LOW_TAIL_STATISTICAL_RARITY"
    # Acidification is a HIGH-side phenomenon, so a low-side RED must not promote
    # a process action, and it must not open any command gate.
    health = evaluate_health(diagnosis)
    assert health["state"] != "PROCESS_ACTION"
    actions = evaluate_operator_actions(diagnosis)
    assert actions["gates"] == {"thermal": "CONDITION_NOT_VERIFIED",
                                "alkali": "CONDITION_NOT_VERIFIED"}
    assert actions["feed"]["target_tpd"] is None


def test_dq_verdicts_are_never_scored_against_the_numeric_bands():
    """Specification sections 3 and 10. The project's own adapter treats every
    non-VALID DQ verdict as not classifiable."""
    fields = [header for header in HEADERS if header != "date"]
    values = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1706, "VFA_TA_B": 0.1715,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    # Section 3 withholds corrupted / faulted / stale readings from the bands.
    for verdict in ("MISSING", "MISSING_ZERO_SENTINEL", "INVALID_PARSE",
                    "INVALID_RANGE", "STALE"):
        statuses = {field: "VALID" for field in fields}
        statuses["pH_A"] = verdict
        diagnosis = diagnose(values, statuses)
        assert diagnosis["classified"]["pH"]["A"]["status"] in ("INVALID", "STALE"), verdict
        assert diagnosis["classified"]["pH"]["A"]["value"] is None, verdict
        signals = {item["variable"]: item
                   for item in evaluate_telemetry_signals(diagnosis, values, statuses)}
        # worstOf: one DATA side and one GREEN side rolls the row up to DATA.
        assert signals["pH"]["status"] == "DATA", verdict
        assert signals["pH"]["status_text"] == "데이터확인", verdict
        assert signals["pH"]["dq_status"] == "UNUSABLE", verdict

    # A statistical SUSPECT is not a corrupted reading. run_data_quality keeps the
    # value in the clean frame on purpose, so it is still classified and the
    # caveat travels alongside it instead of blanking the row.
    for verdict in ("SUSPECT_SPIKE", "SUSPECT_TEMPORAL_OUTLIER",
                    "SUSPECT_FLATLINE", "SUSPECT_CROSS_FIELD"):
        statuses = {field: "VALID" for field in fields}
        statuses["pH_A"] = verdict
        diagnosis = diagnose(values, statuses)
        assert diagnosis["classified"]["pH"]["A"]["status"] == "GREEN", verdict
        assert diagnosis["classified"]["pH"]["A"]["value"] == values["pH_A"], verdict
        signals = {item["variable"]: item
                   for item in evaluate_telemetry_signals(diagnosis, values, statuses)}
        assert signals["pH"]["status"] == "NORMAL", verdict
        assert signals["pH"]["status_text"] == "정상", verdict
        assert signals["pH"]["dq_status"] == "SUSPECT", verdict

    # An unusable field is never masked by a merely suspect sibling.
    statuses = {field: "VALID" for field in fields}
    statuses["VFA_A_mgL"] = "SUSPECT_SPIKE"
    statuses["ALK_A_mgL"] = "MISSING"
    diagnosis = diagnose(values, statuses)
    assert diagnosis["classified"]["VFA_TA"]["A"]["status"] == "INVALID"
    # A derived variable inherits the verdict of the fields it is computed from.
    statuses = {field: "VALID" for field in fields}
    statuses["ALK_B_mgL"] = "INVALID_RANGE"
    diagnosis = diagnose(values, statuses)
    assert diagnosis["classified"]["ALK"]["B"]["status"] == "INVALID"
    assert diagnosis["classified"]["VFA_TA"]["B"]["status"] == "INVALID"


def test_ab_concordance_uses_only_the_project_limits():
    fields = [header for header in HEADERS if header != "date"]
    base = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1706, "VFA_TA_B": 0.1715,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    valid = {field: "VALID" for field in fields}
    signals = {item["variable"]: item for item in evaluate_telemetry_signals(
        diagnose(base, valid), base, valid)}
    assert signals["temperature"]["consistency"] == "CONCORDANT"
    assert signals["temperature"]["ab_limit"] == project_rules.CONCORDANCE["temperature"]["verify"]

    wide = {**base, "temperature_B_C": 38.5 + project_rules.CONCORDANCE["temperature"]["verify"] + 0.1}
    signals = {item["variable"]: item for item in evaluate_telemetry_signals(
        diagnose(wide, valid), wide, valid)}
    assert signals["temperature"]["consistency"] == "DISCORDANT"


def test_operator_actions_come_from_project_conditions_and_gates():
    fields = [header for header in HEADERS if header != "date"]
    values = {
        "temperature_A_C": 35.0, "temperature_B_C": 35.4,   # both outside 36-41
        "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1706, "VFA_TA_B": 0.1715,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    valid = {field: "VALID" for field in fields}
    diagnosis = diagnose(values, valid)
    payload = evaluate_operator_actions(diagnosis)

    assert payload["state"] == "PROCESS_ACTION"
    assert payload["lamp"] == "RED"
    assert payload["control_mode"] == "NO_LIVE_CONTROL"
    thermal = next(a for a in payload["actions"] if a["code"] == "OP-THERMAL-001")
    assert thermal["priority"] == "ACTION"
    assert thermal["reason"] == "THERMAL_ACTION_CANDIDATE"
    assert thermal["rule_source"] == project_rules.RULE_SOURCE
    assert payload["gates"]["thermal"] == "CONDITION_VERIFIED"
    assert payload["gates"]["alkali"] == "CONDITION_NOT_VERIFIED"
    # Thermal is not a feed condition, so the project issues no feed target.
    assert payload["feed"]["target_tpd"] is None
    assert payload["feed"]["q_ref_tpd"] == project_rules.Q_REF_TPD
    assert any("NaHCO3" in advisory for advisory in payload["advisories"])
    # No dose, setpoint or percentage the project does not define may appear.
    text = json.dumps(payload, ensure_ascii=False)
    for invented in ("10%", "500 kg", "2도", "2°C", "setpoint =", "NaHCO3 투입량"):
        assert invented not in text

    # A normal process still produces an instruction, never an empty panel.
    normal = dict(values, temperature_A_C=38.5, temperature_B_C=38.6)
    normal_payload = evaluate_operator_actions(diagnose(normal, valid))
    assert normal_payload["actions"]
    assert normal_payload["actions"][0]["code"] == "OP-DATA-003"   # FAN keeps confidence below HIGH
    assert normal_payload["feed"]["target_tpd"] is None


def test_acid_severe_and_early_reproduce_the_project_feed_formula():
    fields = [header for header in HEADERS if header != "date"]
    base = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
        "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
    }
    valid = {field: "VALID" for field in fields}

    severe = {**base, "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0,
              "VFA_TA_A": 0.50, "VFA_TA_B": 0.52}
    payload = evaluate_operator_actions(diagnose(severe, valid))
    assert payload["actions"][0]["code"] == "OP-ACID-001"
    assert payload["feed"]["target_tpd"] == round(
        project_rules.Q_REF_TPD * project_rules.ACID["severeFeedFactor"], 1)

    early_values = {**base, "VFA_A_mgL": 3200.0, "VFA_B_mgL": 3210.0,
                    "VFA_TA_A": 0.25, "VFA_TA_B": 0.24}
    previous = {**base, "VFA_A_mgL": 2400.0, "VFA_B_mgL": 2410.0,
                "VFA_TA_A": 0.19, "VFA_TA_B": 0.18}
    deltas = build_deltas(early_values, previous)
    payload = evaluate_operator_actions(diagnose(early_values, valid, deltas))
    assert payload["actions"][0]["code"] == "OP-ACID-002"
    assert payload["feed"]["target_tpd"] == round(
        project_rules.Q_REF_TPD * project_rules.ACID["earlyFeedFactor"], 1)
    # Without a previous direct measurement the rule must fail safe.
    assert "ACID_EARLY" not in diagnose(early_values, valid, None)["health"]["code"]


def test_final_demo_workbook_demonstrates_the_project_states():
    """The shipped demo must exercise the project's own conditions, with clean DQ
    on the demonstration dates -- a flagged cell is not classifiable at all and
    would hide the state the date exists to show."""
    workbook = ROOT / "deliverables/BioGuard_2023_Yeongcheon_FINAL_DEMO.xlsx"
    assert workbook.is_file(), "run tools/generate_final_demo_workbook.py"
    raw = parse_workbook(workbook.read_bytes())
    clean, dq, summary = run_data_quality(raw)
    clean = add_derived(clean)

    assert summary["missing"] == 0 and summary["invalid"] == 0
    assert summary["suspect"] < 0.01 * summary["valid"]
    assert clean["CH4_m3d_observed"].notna().all(), "observed CH4 must be continuous"

    fields = [header for header in HEADERS if header != "date"]
    ordered = clean["date"].tolist()

    def statuses_for(stamp):
        rows = dq.loc[dq["date"].eq(stamp)]
        return {str(row.field): str(row.dq_status) for row in rows.itertuples()}

    def values_for(stamp):
        row = clean.loc[clean["date"].eq(stamp)].iloc[0]
        return {key: (None if pd.isna(value) else value) for key, value in row.items()}

    expected = {
        "2023-09-14": ("PROCESS_WATCH", "AMBER", "ACID_WATCH", "OP-ACID-003"),
        "2023-09-15": ("PROCESS_ACTION", "RED", "THERMAL_ACTION_CANDIDATE", "OP-THERMAL-001"),
        "2023-09-16": ("PROCESS_WATCH", "AMBER", "ACID_WATCH", "OP-ACID-003"),
        # RECOVERY stays AMBER: the engine still asks for operator verification.
        "2023-09-17": ("RECOVERY", "AMBER", "RECOVERED", "OP-RECOVERY-001"),
    }
    for text, (state, lamp, code, action_code) in expected.items():
        stamp = pd.Timestamp(text)
        flags = dq.loc[dq["date"].eq(stamp) & ~dq["dq_status"].eq("VALID")]
        assert flags.empty, f"{text} must be DQ clean, got {sorted(set(flags['field']))}"

        position = ordered.index(stamp)
        previous = ordered[position - 1]
        values = values_for(stamp)
        diagnosis = diagnose(
            values, statuses_for(stamp),
            build_deltas(values, values_for(previous)),
            diagnose(values_for(previous), statuses_for(previous))["health"]["state"],
        )
        health = evaluate_health(diagnosis)
        assert health["state"] == state, (text, health["state"])
        assert health["lamp"] == lamp, (text, health["lamp"])
        assert code in health["code"], (text, health["code"])

        actions = evaluate_operator_actions(diagnosis)
        assert any(item["code"] == action_code for item in actions["actions"]), (
            text, [item["code"] for item in actions["actions"]])
        assert actions["control_mode"] == "NO_LIVE_CONTROL"

        signals = {item["variable"]: item for item in
                   evaluate_telemetry_signals(diagnosis, values, statuses_for(stamp))}
        assert len(signals) == 13
        # Every rule variable reports a real classification, never a fallthrough.
        for variable in project_rules.VARIABLE_IDS:
            assert signals[variable]["rule_source"] == project_rules.RULE_SOURCE
        # 09-15 is the thermal demonstration: both channels must be classifiable.
        if text == "2023-09-15":
            assert diagnosis["classified"]["temperature"]["A"]["value"] is not None
            assert diagnosis["classified"]["temperature"]["B"]["value"] is not None
            assert actions["gates"]["thermal"] == "CONDITION_VERIFIED"


def test_forecast_exposes_the_a30_biogas_target_and_purity_matches_it():
    """Gauge 3 needs the model's own biogas prediction, and D+1 purity must be
    derived from the same h1 biogas value rather than a second estimate."""
    client = TestClient(app)
    body = upload_bytes(client, "biogas.xlsx", workbook_bytes(date(2025, 9, 1), 40)).json()
    forecast = client.get(f"/api/v1/datasets/{body['dataset_id']}/forecast").json()
    assert forecast["status"] == "AVAILABLE"

    biogas = forecast["biogas"]["predictions"]
    assert biogas, "the A30 V2 biogas_AB_m3d target must reach the API"
    assert [item["horizon"] for item in biogas] == sorted(item["horizon"] for item in biogas)
    assert biogas[0]["horizon"] == 1
    for item in biogas:
        assert isinstance(item["value"], float) and item["value"] >= 0
        assert item["target_date"]

    methane_h1 = next(i for i in forecast["methane"]["predictions"] if i["horizon"] == 1)
    biogas_h1 = next(i for i in biogas if i["horizon"] == 1)
    assert biogas_h1["target_date"] == methane_h1["target_date"]
    purity = forecast["purity_next_day"]
    if purity and purity.get("value_pct") is not None:
        assert purity["value_pct"] == pytest.approx(
            100.0 * methane_h1["value"] / biogas_h1["value"])
        assert purity["method"] == "DERIVED_FROM_CH4_AND_BIOGAS_FORECAST"


def test_history_declares_whether_observed_ch4_came_from_a_synthesised_input():
    """A demo workbook records which cells it synthesised. The trace must not
    present a value derived from an imputed input as a direct measurement."""
    client = TestClient(app)
    body = upload_bytes(client, "provenance.xlsx", workbook_bytes(date(2025, 9, 1), 30)).json()
    history = client.get(f"/api/v1/datasets/{body['dataset_id']}/history").json()
    assert history["ch4_formula"] == (
        "biogas_A * ch4_purity_A / 100 + biogas_B * ch4_purity_B / 100")
    for row in history["rows"]:
        assert row["ch4_source"] in {
            "DERIVED_FROM_OBSERVED", "DERIVED_FROM_IMPUTED", "UNAVAILABLE"}
        if row["CH4_m3d_observed"] is None:
            assert row["ch4_source"] == "UNAVAILABLE"
    # This fixture has no IMPUTATION_LOG, so nothing may be labelled imputed.
    assert {row["ch4_source"] for row in history["rows"]} == {"DERIVED_FROM_OBSERVED"}

    # The shipped demo workbook does declare synthesised cells.
    demo = ROOT / "deliverables/BioGuard_2023_Yeongcheon_FINAL_DEMO.xlsx"
    log = parse_imputation_log(demo.read_bytes())
    assert len(log) > 0
    assert set(log["field"]) <= set(HEADERS)


def test_only_a_data_confidence_downgrade_is_spared_from_escalating_the_tower():
    """Section 5 of the specification maps every state to a colour. The process
    lamp keeps those colours except for NORMAL_DEGRADED_DATA, which reports "no
    process-action evidence" and must not read as a process alarm."""
    fields = [header for header in HEADERS if header != "date"]
    base = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1706, "VFA_TA_B": 0.1715,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    valid = {field: "VALID" for field in fields}

    calm = evaluate_health(diagnose(base, valid))
    assert calm["state"] == "NORMAL_DEGRADED_DATA"
    assert (calm["lamp"], calm["combined_lamp"]) == ("GREEN", "AMBER")

    # Every genuine process condition still escalates exactly as specified.
    watch = evaluate_health(diagnose({**base, "VFA_TA_A": 0.23, "VFA_TA_B": 0.24}, valid))
    assert watch["state"] == "PROCESS_WATCH" and watch["lamp"] == "AMBER"
    action = evaluate_health(
        diagnose({**base, "temperature_A_C": 35.0, "temperature_B_C": 35.4}, valid))
    assert action["state"] == "PROCESS_ACTION" and action["lamp"] == "RED"
    suspect = evaluate_health(diagnose({**base, "VFA_A_mgL": 3700.0}, valid))
    assert suspect["state"] == "DATA_SUSPECT" and suspect["lamp"] == "AMBER"
    for state, lamp in project_rules.STATE_TO_LAMP.items():
        if state != "NORMAL_DEGRADED_DATA":
            assert lamp == project_rules.STATE_TO_LAMP[state]


def test_operator_recommendations_are_specific_and_never_command_equipment():
    """Severity comes from the project; reference practice only supplies the
    response. Nothing may issue a physical command or invent a dose/setpoint."""
    client = TestClient(app)
    body = upload_bytes(client, "recommend.xlsx", workbook_bytes(date(2025, 9, 1), 40)).json()
    row = client.get(
        f"/api/v1/datasets/{body['dataset_id']}/rows/{body['forecast_origin_date']}"
    ).json()
    payload = row["operator_actions"]
    recommendations = payload["recommendations"]
    assert recommendations, "the panel is never empty"
    assert payload["control_mode"] == "NO_LIVE_CONTROL"

    priorities = [item["priority"] for item in recommendations]
    order = {"ACTION": 0, "WATCH": 1, "DATA_REVIEW": 2, "NORMAL": 3}
    assert priorities == sorted(priorities, key=lambda value: order[value])
    for item in recommendations:
        assert item["priority"] in order
        assert item["code"].startswith("REC-")
        assert item["title"] and item["actions"] and item["recheck"]
        assert 1 <= len(item["actions"]) <= 5, item["code"]
        assert set(item["source"]) <= {"PROJECT", "ENGINEERING_REFERENCE"}

    text = json.dumps(payload, ensure_ascii=False)
    # No physical command, and no dose / setpoint the project does not define.
    for forbidden in ("밸브", "펌프를 기동", "heater start", "자동 기동", "자동 개폐",
                      "kg 투입", "NaHCO3 투입량", "setpoint 을", "rpm"):
        assert forbidden not in text, forbidden
    # The one feed figure allowed is the project's own factor times Q_ref.
    if payload["feed"]["target_tpd"] is not None:
        assert payload["feed"]["target_tpd"] in {
            round(project_rules.Q_REF_TPD * project_rules.ACID["severeFeedFactor"], 1),
            round(project_rules.Q_REF_TPD * project_rules.ACID["earlyFeedFactor"], 1),
        }


def test_a_trend_alone_never_produces_an_action_recommendation():
    """Section 39: a forecast or a trend is supporting information. Only a
    project process condition may raise ACTION."""
    fields = [header for header in HEADERS if header != "date"]
    values = {
        "temperature_A_C": 38.5, "temperature_B_C": 38.6, "pH_A": 7.9, "pH_B": 7.92,
        "VFA_A_mgL": 2900.0, "VFA_B_mgL": 2950.0, "ALK_A_mgL": 17000.0, "ALK_B_mgL": 17200.0,
        "VFA_TA_A": 0.1706, "VFA_TA_B": 0.1715,
        "TAN_A_mgL": 4000.0, "TAN_B_mgL": 4050.0, "TS_A_pct": 2.6, "TS_B_pct": 2.6,
        "VS_A_pct": 1.1, "VS_B_pct": 1.1, "CODcr_A_mgL": 25000.0, "CODcr_B_mgL": 25000.0,
    }
    valid = {field: "VALID" for field in fields}
    diagnosis = diagnose(values, valid)
    evidence = build_evidence(diagnosis, valid)

    # A collapse in gas output with no project condition present.
    history = [
        {"date": f"2025-09-{day:02d}", "feed_AB_tpd": 180.0, "biogas_AB_m3d": 9900.0,
         "CH4_m3d_observed": 6200.0, "ch4_purity_A_pct": 62.0, "ch4_purity_B_pct": 63.0}
        for day in range(1, 10)
    ]
    history.append({"date": "2025-09-10", "feed_AB_tpd": 180.0, "biogas_AB_m3d": 4000.0,
                    "CH4_m3d_observed": 2400.0, "ch4_purity_A_pct": 48.0,
                    "ch4_purity_B_pct": 49.0})
    trend = build_trend(history, "2025-09-10")
    assert trend["fields"]["biogas_AB_m3d"]["notable"] is True
    assert trend["fields"]["biogas_AB_m3d"]["direction"] == "DOWN"

    recommendations = build_recommendations(diagnosis, evidence, trend, None)
    assert recommendations
    assert all(item["priority"] != "ACTION" for item in recommendations), (
        "a trend on its own must not escalate to ACTION")
    assert any(item["code"] == "REC-GAS-001" for item in recommendations)


def test_console_assets_are_revalidated_so_no_stale_document_or_adapter_is_served():
    """A cached document paired with a fresh adapter (or the reverse) renders a
    console that matches neither build, so both must revalidate."""
    client = TestClient(app)
    for path in ("/", "/datasets/anything", "/static/bioguard_v2.js"):
        response = client.get(path)
        assert response.status_code == 200, path
        assert "no-cache" in response.headers.get("cache-control", ""), path

    # The served document must be the current build, not a stale deliverable.
    html = client.get("/").text
    for title in ("메탄생성량 예측", "메탄순도 예측",
                  "총 바이오가스 생산량 예측", "소화조 내부 온도 / pH"):
        assert title in html, title
    live = html.split('<script type="application/x-bioguard-legacy-reference"')[0]
    assert 'id="gaugeVfaTaValue"' not in live


def test_final_demo_workbook_keeps_its_documentation_sheets():
    workbook = ROOT / "deliverables/BioGuard_2023_Yeongcheon_FINAL_DEMO.xlsx"
    names = load_workbook(workbook, read_only=True).sheetnames
    assert names[0] == "BIOGUARD_INPUT"
    assert set(names) == {
        "BIOGUARD_INPUT", "IMPUTATION_LOG", "REFERENCE_STATS", "SIGNAL_TEST_GUIDE", "README",
    }
    log = pd.read_excel(workbook, sheet_name="IMPUTATION_LOG")
    assert list(log.columns) == [
        "date", "field", "raw_value", "replacement_value", "reason", "method", "provenance",
    ]
    assert set(log["provenance"]) == {"SYNTHETIC_DEMO_REPLACEMENT"}
    assert "SIGNAL_DEMONSTRATION_FIXTURE" in set(log["reason"])


def test_row_api_returns_the_full_project_signal_and_action_payload():
    client = TestClient(app)
    body = upload_bytes(client, "signals.xlsx", workbook_bytes(date(2025, 9, 1), 30)).json()
    row = client.get(
        f"/api/v1/datasets/{body['dataset_id']}/rows/{body['forecast_origin_date']}"
    ).json()

    signals = row["telemetry_signals"]
    assert len(signals) == 13
    assert [item["variable"] for item in signals][:10] == project_rules.VARIABLE_IDS
    assert {item["status"] for item in signals} <= {
        "NORMAL", "WATCH", "ACTION", "DATA", "DATA_NORMAL", "DATA_WATCH", "INFO"
    }
    for item in signals:
        assert item["rule_source"]
        assert item["status_text"]
        assert "reason" in item

    health = row["health"]
    assert health["state"] in project_rules.STATE_TO_LAMP
    assert health["lamp"] == project_rules.STATE_TO_LAMP[health["state"]]
    assert health["data_confidence"] in {"HIGH", "MEDIUM", "LOW"}
    assert health["rule_source"] == project_rules.RULE_SOURCE

    actions = row["operator_actions"]
    assert actions["actions"], "the operator panel is never empty"
    assert actions["control_mode"] == "NO_LIVE_CONTROL"
    assert set(actions["gates"]) == {"thermal", "alkali"}
    for action in actions["actions"]:
        assert action["code"] and action["title"] and action["recommendation"]
        assert action["priority"] in {"ACTION", "WATCH", "NORMAL"}
        assert action["verdict"], "each action states its verdict"
        assert action["rule_source"] == project_rules.RULE_SOURCE
    assert actions["feed"]["evidence"], "the feed line states why"

    # Process health and data confidence are reported separately.
    assert health["process_state"] in {"NORMAL", "PROCESS_WATCH", "PROCESS_ACTION", "DATA_SUSPECT"}
    assert health["process_text"]
    evidence = health["evidence"]
    assert set(evidence) >= {"process", "data_quality", "suspect_fields",
                             "data_confidence", "process_normal"}
    for item in evidence["process"]:
        assert item["code"] and item["level"]
        for reading in item["readings"]:
            assert reading["variable"] in project_rules.THRESHOLDS
            assert reading["green_range"] == [
                project_rules.THRESHOLDS[reading["variable"]]["lowGreen"],
                project_rules.THRESHOLDS[reading["variable"]]["highGreen"],
            ]
            assert reading["rule_source"] == project_rules.RULE_SOURCE
    # A process condition and a degraded channel are never merged.
    process_codes = {item["code"] for item in evidence["process"]}
    data_variables = {item["variable"] for item in evidence["data_quality"]}
    assert not (process_codes & data_variables)
