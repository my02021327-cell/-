# -*- coding: utf-8 -*-
"""조내 VS · 유출 VS 기준 SRT 산정 방안의 평가.

■ 사용자 지시
   · 2018~2022 운영일지 원본(엑셀 5개)을 참조한다. 2023년은 온도 붕괴·VFA 붕괴가 있어 제외.
   · master CSV와 값이 다르면 **master를 우선 신뢰**한다.
   · SRT를 '조내 VS'와 '유출 VS' 기준으로 구하는 방안을 평가한다.

■ 평가 대상 네 가지 경로
   V1. 운영일지 내장 방식 — 조내 유기물량 ÷ 유출 유기물량
   V2. 실측 유출 고형물 기준 — 탈수 케이크 + 탈리액으로 실제 빠져나가는 VS를 계산
   V3. VS 파괴 기반 1차 반응 — VS_in/VS_out − 1 = k_d·SRT
   V4. VS 동적 응답 — TS로는 실패했던 시상수 추정을 VS·VS/TS로 재시도
"""
import glob
import json
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

UP = "/root/.claude/uploads/1c0f73df-ed3f-50b0-a104-6fa33713e03c"
OUT = {}

# ================================================================ 1. 원본 추출
WANT = {  # (블록 키워드, 항목명) → 컬럼명
    ("반입", "반입량"): "wb_intake_tpd",
    ("유기산화조", "TS(%)"): "wb_acid_TS",
    ("유기산화조", "VS(%)"): "wb_acid_VS",
    ("#A", "투입량"): "wb_feed_A", ("#B", "투입량"): "wb_feed_B",
    ("#A", "TS(%)"): "wb_dig_TS_A", ("#B", "TS(%)"): "wb_dig_TS_B",
    ("#A", "VS(%)"): "wb_dig_VS_A", ("#B", "VS(%)"): "wb_dig_VS_B",
    ("#A", "온도(중)"): "wb_T_A", ("#B", "온도(중)"): "wb_T_B",
    ("#A", "VFA"): "wb_VFA_A", ("#B", "VFA"): "wb_VFA_B",
    ("#A", "Alk"): "wb_ALK_A", ("#B", "Alk"): "wb_ALK_B",
    ("#A", "바이오가스발생량"): "wb_gas_A", ("#B", "바이오가스발생량"): "wb_gas_B",
    ("#A", "메탄함량"): "wb_CH4_A", ("#B", "메탄함량"): "wb_CH4_B",
    ("탈리액", "탈수량"): "wb_filtrate_tpd",
    ("탈리액", "SS"): "wb_filtrate_SS",
    ("탈리액", "TS(%)"): "wb_filtrate_TS",
    ("탈리액", "VS(%)"): "wb_filtrate_VS",
    ("소화조 탈수슬러지", "발생량"): "wb_cake_dig_tpd",
    ("소화조 탈수슬러지", "TS(%)"): "wb_cake_dig_TS",
    ("혼합슬러지 탈수슬러지", "발생량"): "wb_cake_mix_tpd",
    ("혼합슬러지 탈수슬러지", "TS(%)"): "wb_cake_mix_TS",
    ("MVR", "VS(%)"): "wb_mvr_VS",
}


def read_year(f):
    """read_only 시트는 cell() 임의접근이 극도로 느리다. 한 번만 스트리밍해 메모리에 올린다."""
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb["Data"]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    ncol = max(len(r) for r in grid)
    g = lambda r, c: (grid[r - 1][c - 1] if r - 1 < len(grid) and c - 1 < len(grid[r - 1])
                      else None)

    # 병합 해제된 상위 헤더를 왼쪽으로 전파해 컬럼별 블록명 복원
    blocks = []
    for r in (2, 3):
        cur, row = "", []
        for c in range(1, ncol + 1):
            v = g(r, c)
            if v not in (None, ""):
                cur = str(v)
            row.append(cur)
        blocks.append(row)
    colmap = {}
    for c in range(1, ncol + 1):
        h4 = g(4, c)
        if not h4:
            continue
        h4 = str(h4).strip()
        blk = blocks[0][c - 1] + "/" + blocks[1][c - 1]
        for (bk, it), name in WANT.items():
            if name not in colmap and it == h4 and bk in blk:
                colmap[name] = c

    rows = []
    for r in range(6, len(grid) + 1):
        d = g(r, 1)
        if not hasattr(d, "year"):
            continue
        rec = {"date": pd.Timestamp(d).normalize()}
        for name, c in colmap.items():
            v = g(r, c)
            rec[name] = v if isinstance(v, (int, float)) else np.nan
        rows.append(rec)
    return pd.DataFrame(rows).set_index("date"), colmap


files, seen = {}, {}
for f in sorted(glob.glob(f"{UP}/*BGP*.xlsx")):
    y = re.search(r"/\w+-(\d\d)", f).group(1)
    seen.setdefault(y, f)
frames = []
print("=== 1. 운영일지 원본 추출 ===")
for y, f in sorted(seen.items()):
    df, cm = read_year(f)
    frames.append(df)
    print(f"  20{y}년  {len(df):4d}행, 매칭 컬럼 {len(cm)}/{len(WANT)}개"
          f"{'   ★미매칭: ' + ','.join(sorted(set(WANT.values()) - set(cm))) if len(cm) < len(WANT) else ''}")
W = pd.concat(frames).sort_index()
W = W[~W.index.duplicated(keep="first")]
W = W.loc["2018-01-01":"2022-12-31"]
print(f"  통합 {len(W)}행 ({W.index.min().date()} ~ {W.index.max().date()})")

# ================================================================ 2. master 대조
M = pd.read_csv("data/영천BGP_MASTER_2018-2023.csv", parse_dates=["date"]).set_index("date")
M = M.loc["2018-01-01":"2022-12-31"]
PAIRS = [("wb_feed_A", "feed_A_tpd"), ("wb_feed_B", "feed_B_tpd"),
         ("wb_acid_TS", "acid_TS_pct"), ("wb_acid_VS", "acid_VS_pct"),
         ("wb_dig_TS_A", "dig_TS_A_pct"), ("wb_dig_VS_A", "dig_VS_A_pct"),
         ("wb_dig_TS_B", "dig_TS_B_pct"), ("wb_dig_VS_B", "dig_VS_B_pct"),
         ("wb_VFA_A", "VFA_A_mgL"), ("wb_ALK_A", "ALK_A_mgL"),
         ("wb_gas_A", "biogas_A_m3d"), ("wb_gas_B", "biogas_B_m3d"),
         ("wb_intake_tpd", "intake_total_tpd")]
print("\n=== 2. master 대조 (불일치 시 master 우선) ===")
comp = []
for w, m in PAIRS:
    if w not in W or m not in M:
        continue
    d = pd.concat([W[w].rename("w"), M[m].rename("m")], axis=1).dropna()
    if not len(d):
        continue
    rel = np.abs(d.w - d.m) / d.m.replace(0, np.nan)
    bad = int((rel > 0.01).sum())
    comp.append({"col": m, "n": int(len(d)), "mismatch_1pct": bad,
                 "mismatch_pct": round(100 * bad / len(d), 1),
                 "median_rel_diff_pct": round(float(rel.median() * 100), 3)})
    print(f"  {m:18s} n={len(d):4d}  1%초과 불일치 {bad:4d}건 ({100*bad/len(d):5.1f}%)  "
          f"중앙 상대차 {rel.median()*100:6.3f}%")
OUT["master_check"] = comp

# master 우선: 겹치는 항목은 master 값으로 덮고, master에 없는 항목만 운영일지에서 가져온다
NEW = ["wb_filtrate_tpd", "wb_filtrate_TS", "wb_filtrate_VS", "wb_filtrate_SS",
       "wb_cake_dig_tpd", "wb_cake_dig_TS", "wb_cake_mix_tpd", "wb_cake_mix_TS", "wb_mvr_VS"]
D = pd.DataFrame(index=pd.date_range("2018-01-01", "2022-12-31", freq="D"))
D["feed"] = M.feed_AB_tpd
D["VS_in"] = M.acid_VS_pct                      # 유기산화조(소화조 투입) VS%  ← master 우선
D["TS_in"] = M.acid_TS_pct
D["VS_r"] = M[["dig_VS_A_pct", "dig_VS_B_pct"]].mean(axis=1)   # 조내 VS%
D["TS_r"] = M[["dig_TS_A_pct", "dig_TS_B_pct"]].mean(axis=1)
D["CH4"] = M.CH4_m3d
for c in NEW:
    D[c] = W[c] if c in W else np.nan
print(f"  master에 없어 운영일지에서만 가져온 항목: {', '.join(c for c in NEW if D[c].notna().any())}")
print(f"  → 탈리액 VS 관측 {int(D.wb_filtrate_VS.notna().sum())}일, "
      f"소화조 탈수케이크 관측 {int(D.wb_cake_dig_tpd.notna().sum())}일")

V = 8000.0  # 설계 용적 (유효 용적은 미지 — §2.8 참조)

# ================================================================ V1. 운영일지 내장 방식
print("\n=== V1. 운영일지 내장 방식 — 조내 유기물량 ÷ 유출 유기물량 ===")
inv_VS = V * D.VS_r / 100 * 1000                      # kg VS  (= V × 조내VS%)
out_VS_wb = D.feed * D.VS_r / 100 * 1000              # kg VS/d (= Q_in × 조내VS%)
srt_v1 = (inv_VS / out_VS_wb).replace([np.inf, -np.inf], np.nan)
hrt = (V / D.feed).replace([np.inf, -np.inf], np.nan)
same = float(np.nanmax(np.abs(srt_v1 - hrt)))
OUT["V1"] = {"srt_median": round(float(srt_v1.median()), 2),
             "hrt_median": round(float(hrt.median()), 2),
             "max_abs_diff": round(same, 9), "identical": bool(same < 1e-6),
             "n": int(srt_v1.notna().sum())}
print(f"  SRT(VS) 중앙값 {srt_v1.median():.2f}일 / HRT 중앙값 {hrt.median():.2f}일")
print(f"  두 계열의 최대 절대차 {same:.2e}일  →  {'완전히 동일' if same < 1e-6 else '다름'}")
print("  이유: 운영일지의 '유출 유기물량'은 실측 유출 VS가 아니라")
print("        소화조 투입량 × 조내 VS% 로 정의돼 있다. VS%가 분자·분모에서 약분되어")
print("        SRT = (V×VS)/(Q×VS) = V/Q = HRT 가 되며, VS 정보가 사라진다.")
# 2018 운영일지 원본 셀과 대조
try:
    wb18 = openpyxl.load_workbook(seen["18"], read_only=True, data_only=True)
    ws = wb18["Data 계산"]
    G = [list(r) for r in ws.iter_rows(values_only=True)]
    wb18.close()
    gg = lambda r, c: (G[r - 1][c - 1] if r - 1 < len(G) and c - 1 < len(G[r - 1]) else None)
    chk = []
    for r in range(7, min(400, len(G) + 1)):
        dt, outflow = gg(r, 45), gg(r, 45)
        dt = gg(r, 1)
        outflow, q, vsr = gg(r, 45), gg(r, 5), gg(r, 32)
        if not hasattr(dt, "year"):
            continue
        if all(isinstance(v, (int, float)) for v in (outflow, q, vsr)) and vsr:
            chk.append(abs(outflow - q * vsr / 100 * 1000))
    if chk:
        print(f"  [원본 검증] 2018년 {len(chk)}행에서 '유출 유기물량 = 투입량x조내VS%' 최대 오차 {max(chk):.2e} kg")
        OUT["V1"]["workbook_formula_verified"] = bool(max(chk) < 1.0)
        OUT["V1"]["workbook_formula_n"] = len(chk)
except Exception as e:
    print(f"  [원본 검증] 생략 ({type(e).__name__}: {e})")

# ================================================================ V2. 실측 유출 고형물
print("\n=== V2. 실측 유출 고형물 기준 — SRT = V·X_r / (Q_w·X_w + Q_e·X_e) ===")
have = {c: int(D[c].notna().sum()) for c in NEW}
print("  운영일지 유출 계열 데이터 충전 현황:")
for c, n in have.items():
    print(f"    {c:20s} {n:5d}일" + ("   ★전무" if n == 0 else ""))
# 탈리액 SS (mg/L) — 운영일지에 유일하게 채워져 있는 유출 고형물 지표
ss = W["wb_filtrate_SS"] if "wb_filtrate_SS" in W else pd.Series(dtype=float)
OUT["V2"] = {
    "constructible": False,
    "have": have,
    "filtrate_SS_n": int(ss.notna().sum()) if len(ss) else 0,
    "filtrate_SS_median_mgL": (round(float(ss.median()), 1) if len(ss) and ss.notna().any() else None),
    "filtrate_flow_median_tpd": (round(float(D.wb_filtrate_tpd.median()), 1)
                                 if D.wb_filtrate_tpd.notna().any() else None),
    "reason": ("탈리액 VS(%) 컬럼은 운영일지 서식에 존재하나 전 기간 미기입. "
               "탈수 케이크의 TS(%)도 2019년 이후 서식에서 사라졌다. "
               "게다가 탈리액은 '소화조 + 슬러지저류조' 혼합 스트림이라 "
               "값이 있어도 소화조 몫으로 귀속할 수 없다.")}
print(f"  → 구성 불가. 이유: 유출 VS 실측치가 없다.")
print(f"     탈리액 VS(%)는 서식에만 존재하고 전 기간 미기입,")
print(f"     탈수 케이크 TS(%)는 2019년 이후 서식에서 삭제,")
print(f"     탈리액은 소화조+슬러지저류조 혼합이라 소화조 몫 귀속 불가.")
if OUT["V2"]["filtrate_SS_median_mgL"]:
    q = OUT["V2"]["filtrate_flow_median_tpd"]
    c_ = OUT["V2"]["filtrate_SS_median_mgL"]
    esc = q * c_ / 1000 if q else None
    OUT["V2"]["filtrate_solids_kgd"] = round(esc, 1) if esc else None
    OUT["V2"]["pct_of_inventory"] = (round(100 * esc / float(inv_VS.median()), 2)
                                     if esc else None)
    print(f"     참고: 유일하게 채워진 탈리액 SS {c_:.0f} mg/L × 유량 {q:.0f} t/d "
          f"= {esc:.0f} kg/d (조내 재고의 {100*esc/float(inv_VS.median()):.2f}%/일)")
    print(f"           이 값만으로 SRT를 구하면 {float(inv_VS.median())/esc:.0f}일이 나오지만,")
    print(f"           탈수 케이크로 빠지는 고형물이 빠져 있어 **상한**일 뿐 SRT가 아니다.")
    OUT["V2"]["srt_upper_bound_filtrate_only"] = round(float(inv_VS.median()) / esc, 1)

# ================================================================ V3. VS 파괴 1차 반응
print("\n=== V3. VS 파괴 기반 — VS_in/VS_out − 1 = k_d·SRT ===")
d3 = pd.concat([(D.feed * D.VS_in / 100 * 1000).rename("Lin"),
                D.VS_r.rename("Xr"), D.feed.rename("Q"), D.CH4.rename("CH4")],
               axis=1).replace(0, np.nan).dropna()
d3 = d3[(d3.Xr > 0.3) & (d3.Q > 50)]
d3["VSin_pct"] = d3.Lin / (d3.Q * 1000) * 100
d3["ratio"] = d3.VSin_pct / d3.Xr
red = 1 - 1 / d3.ratio
kd_srt = d3.ratio - 1                                  # = k_d · SRT
# (X_in/X − 1) 를 1/Q 에 회귀 → 기울기 = k_d·V
x = 1 / d3.Q.values
A = np.c_[np.ones(len(x)), x]
b, *_ = np.linalg.lstsq(A, kd_srt.values, rcond=None)
pred = A @ b
r2 = float(1 - ((kd_srt.values - pred) ** 2).sum() / ((kd_srt.values - kd_srt.mean()) ** 2).sum())
kdV = float(b[1])
OUT["V3"] = {"n": int(len(d3)), "VS_removal_pct": round(float(red.median() * 100), 1),
             "kd_times_SRT": round(float(kd_srt.median()), 3),
             "regression_R2": round(r2, 4), "kd_times_V": round(kdV, 1),
             "intercept": round(float(b[0]), 3),
             "kd_if_V8000": round(kdV / V, 4),
             "srt_if_kd_0p10": round(float(kd_srt.median()) / 0.10, 1),
             "srt_if_kd_0p05": round(float(kd_srt.median()) / 0.05, 1),
             "srt_if_kd_0p20": round(float(kd_srt.median()) / 0.20, 1)}
print(f"  n={len(d3)}일, VS 제거율 중앙 {red.median()*100:.1f}%  →  k_d·SRT = {kd_srt.median():.3f}")
print(f"  (X_in/X − 1) ~ 1/Q 회귀: 기울기 k_d·V = {kdV:.1f}, 절편 {b[0]:.3f}, R² {r2:.4f}")
print(f"  → k_d 와 SRT 는 **곱으로만** 나타난다. k_d를 모르면 SRT가 결정되지 않는다:")
for k_ in (0.05, 0.10, 0.20):
    print(f"      k_d={k_:.2f}/d 이면 SRT = {kd_srt.median()/k_:5.1f}일")

# 메탄 교차검증 — VS 파괴량당 메탄
dest = (d3.Lin - d3.Q * d3.Xr / 100 * 1000)
ymeth = (d3.CH4 / dest).replace([np.inf, -np.inf], np.nan).dropna()
OUT["V3"]["CH4_per_VS_destroyed"] = round(float(ymeth.median()), 3)
print(f"  [교차검증] 파괴 VS 1kg당 메탄 {ymeth.median():.3f} ㎥ "
      f"(이론 상한 0.50~0.55 대비 {ymeth.median()/0.52*100:.0f}%)")

# ================================================================ V4. VS 동적 응답
print("\n=== V4. VS·VS/TS 동적 응답 — TS로 실패했던 시상수 추정 재시도 ===")


def ewm(s, tau, K=400):
    w = np.exp(-np.arange(K + 1) / tau); w /= w.sum()
    v = np.nan_to_num(s.values); o = np.zeros(len(v))
    for i, wv in enumerate(w):
        o[i:] += wv * v[: len(v) - i]
    return pd.Series(o, index=s.index)


TAUS = [1, 2, 3, 5, 8, 12, 16, 20, 25, 30, 40, 50, 60, 80, 100, 150, 200]
drv = (D.feed * D.VS_in / 100).interpolate(limit=3)
res4 = {}
for nm, tgt in [("소화액 VS", D.VS_r), ("소화액 VS/TS", (D.VS_r / D.TS_r)),
                ("소화액 TS(대조)", D.TS_r)]:
    rows = []
    for t in TAUS:
        dd = pd.concat([ewm(drv, t).rename("x"), tgt.rename("y")], axis=1).dropna()
        if len(dd) < 50:
            continue
        Am = np.c_[np.ones(len(dd)), dd.x.values]
        bb = np.linalg.lstsq(Am, dd.y.values, rcond=None)[0]
        p = Am @ bb
        rows.append({"tau": t, "R2": round(float(1 - ((dd.y - p) ** 2).sum()
                                               / ((dd.y - dd.y.mean()) ** 2).sum()), 4)})
    bst = max(rows, key=lambda r: r["R2"])
    cv = float(tgt.std() / tgt.mean() * 100)
    res4[nm] = {"scan": rows, "best": bst, "cv_pct": round(cv, 1),
                "at_edge": bool(bst["tau"] == max(TAUS))}
    print(f"  [{nm:14s}] 변동계수 {cv:4.1f}%  최적 시상수 {bst['tau']:3d}일  R² {bst['R2']:.4f}"
          f"{'   ★격자 상한 — 신호 없음' if bst['tau'] == max(TAUS) else ''}")
# 내부 최적점의 안정성 — 30일 블록 부트스트랩 200회
print("  [부트스트랩] 최적 시상수의 안정성 (30일 블록 200회 재표본)")
rng = np.random.default_rng(0)
for nm, tgt in [("소화액 VS", D.VS_r), ("소화액 VS/TS", (D.VS_r / D.TS_r))]:
    dd = pd.concat([drv.rename("x0"), tgt.rename("y")], axis=1).dropna()
    idx = np.arange(len(dd)); blocks = [idx[i:i + 30] for i in range(0, len(idx), 30)]
    best = []
    for _ in range(200):
        pick = np.concatenate([blocks[i] for i in rng.integers(0, len(blocks), len(blocks))])
        sub = dd.iloc[pick]
        rs = []
        for t in TAUS:
            xx = ewm(pd.Series(sub.x0.values, index=pd.RangeIndex(len(sub))), t).values
            yy = sub.y.values
            Am = np.c_[np.ones(len(xx)), xx]
            bb = np.linalg.lstsq(Am, yy, rcond=None)[0]
            pp = Am @ bb
            rs.append((float(1 - ((yy - pp) ** 2).sum() / ((yy - yy.mean()) ** 2).sum()), t))
        best.append(max(rs)[1])
    best = np.array(best)
    lo, hi = np.percentile(best, [2.5, 97.5])
    edge = float((best == max(TAUS)).mean() * 100)
    res4[nm]["boot_median_tau"] = float(np.median(best))
    res4[nm]["boot_ci"] = [float(lo), float(hi)]
    res4[nm]["boot_pct_at_edge"] = round(edge, 1)
    print(f"    {nm:14s} 최적 시상수 중앙 {np.median(best):5.1f}일, 95% 구간 {lo:.0f}~{hi:.0f}일, "
          f"격자상한 채택률 {edge:.0f}%")

OUT["V4"] = res4

OUT["meta"] = {"period": "2018-01-01~2022-12-31", "n_days": int(len(D)),
               "V_design_m3": V, "note_2023": "온도·VFA 붕괴로 제외",
               "master_priority": True}
Path("outputs").mkdir(exist_ok=True)
json.dump(OUT, open("outputs/vs_srt_results.json", "w", encoding="utf-8"), ensure_ascii=False)
print("\n저장: outputs/vs_srt_results.json")
